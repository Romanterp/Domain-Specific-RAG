"""
Extract plain text from PDFs and Word documents.

- PDFs: PyMuPDF (fitz) primary, pdfplumber fallback
- DOCX: python-docx
- DOC: fallback to reading raw bytes for any ASCII text

Outputs .txt files alongside metadata about page count and word count.
"""

import json
import logging
import re
from pathlib import Path

import fitz  # PyMuPDF
import docx

DATA_DIR = Path(__file__).resolve().parent.parent / "data"
PDF_DIR = DATA_DIR / "pdfs"
TEXT_DIR = DATA_DIR / "texts"
CATALOG_PATH = DATA_DIR / "catalog.json"
DOC_EXTENSIONS = {".pdf", ".docx", ".doc"}

logging.basicConfig(level=logging.INFO, format="%(asctime)s [%(levelname)s] %(message)s")
log = logging.getLogger(__name__)


def extract_text_from_pdf(pdf_path: Path) -> dict:
    """Extract text from a PDF and return metadata."""
    result = {
        "text": "",
        "pages": [],  # list of {"page": int, "text": str}
        "page_count": 0,
        "word_count": 0,
        "pdf_metadata": {},
    }

    try:
        doc = fitz.open(str(pdf_path))
        result["page_count"] = len(doc)

        # Extract PDF-level metadata
        meta = doc.metadata
        if meta:
            result["pdf_metadata"] = {
                k: v for k, v in meta.items()
                if v and isinstance(v, str) and v.strip()
            }

        # Extract text page by page, skip blank pages
        page_texts = []
        for page_num in range(len(doc)):
            page = doc.load_page(page_num)
            text = page.get_text("text").strip()
            if text:
                page_texts.append({"page": page_num + 1, "text": text})

        result["pages"] = page_texts
        full_text = "\n\n".join(p["text"] for p in page_texts)
        full_text = re.sub(r"\n{4,}", "\n\n\n", full_text)
        result["text"] = full_text
        result["word_count"] = len(full_text.split())

        doc.close()

    except Exception as e:
        log.error(f"PyMuPDF failed for {pdf_path.name}: {e}")
        # Try pdfplumber as fallback
        try:
            import pdfplumber
            with pdfplumber.open(str(pdf_path)) as pdf:
                result["page_count"] = len(pdf.pages)
                page_texts = []
                for page_num, page in enumerate(pdf.pages):
                    text = (page.extract_text() or "").strip()
                    if text:
                        page_texts.append({"page": page_num + 1, "text": text})
                result["pages"] = page_texts
                full_text = "\n\n".join(p["text"] for p in page_texts)
                full_text = re.sub(r"\n{4,}", "\n\n\n", full_text)
                result["text"] = full_text
                result["word_count"] = len(full_text.split())
        except Exception as e2:
            log.error(f"pdfplumber also failed for {pdf_path.name}: {e2}")

    return result


def extract_text_from_docx(docx_path: Path) -> dict:
    """Extract text from a .docx file."""
    result = {
        "text": "",
        "page_count": 0,
        "word_count": 0,
        "pdf_metadata": {},
    }

    try:
        doc = docx.Document(str(docx_path))

        # Extract core properties as metadata
        props = doc.core_properties
        meta = {}
        if props.title:
            meta["title"] = props.title
        if props.author:
            meta["author"] = props.author
        if props.subject:
            meta["subject"] = props.subject
        if props.created:
            meta["creationDate"] = str(props.created)
        if meta:
            result["pdf_metadata"] = meta

        # Extract all paragraph text
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]

        # Also extract text from tables
        for table in doc.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    paragraphs.append(" | ".join(cells))

        full_text = "\n\n".join(paragraphs)
        result["text"] = full_text
        result["word_count"] = len(full_text.split())

    except Exception as e:
        log.error(f"python-docx failed for {docx_path.name}: {e}")

    return result


def extract_text_from_doc(doc_path: Path) -> dict:
    """Best-effort text extraction from legacy .doc files.

    Scans the binary for ASCII text runs. Lossy but better than nothing.
    """
    result = {
        "text": "",
        "page_count": 0,
        "word_count": 0,
        "pdf_metadata": {},
    }

    try:
        raw = doc_path.read_bytes()
        # Extract printable ASCII runs of 20+ chars
        import re
        text_runs = re.findall(rb"[\x20-\x7E\r\n\t]{20,}", raw)
        text = "\n".join(run.decode("ascii", errors="ignore") for run in text_runs)

        # Filter out binary noise (runs that are mostly non-letter)
        lines = []
        for line in text.split("\n"):
            letters = sum(1 for c in line if c.isalpha())
            if letters > len(line) * 0.3:
                lines.append(line.strip())

        full_text = "\n".join(lines)
        result["text"] = full_text
        result["word_count"] = len(full_text.split())

    except Exception as e:
        log.error(f"Failed to extract text from {doc_path.name}: {e}")

    return result


def extract_text(file_path: Path) -> dict:
    """Route to the correct extractor based on file extension."""
    ext = file_path.suffix.lower()
    if ext == ".pdf":
        return extract_text_from_pdf(file_path)
    elif ext == ".docx":
        return extract_text_from_docx(file_path)
    elif ext == ".doc":
        return extract_text_from_doc(file_path)
    else:
        log.warning(f"Unsupported file type: {file_path.name}")
        return {"text": "", "page_count": 0, "word_count": 0, "pdf_metadata": {}}


def extract_all_texts(force: bool = False):
    """Extract text from all downloaded documents."""
    TEXT_DIR.mkdir(parents=True, exist_ok=True)

    doc_files = sorted(
        f for f in PDF_DIR.iterdir()
        if f.suffix.lower() in DOC_EXTENSIONS and f.stat().st_size > 500
    )
    log.info(f"Found {len(doc_files)} documents to process.")

    # Load catalog for updating
    catalog = {}
    if CATALOG_PATH.exists():
        with open(CATALOG_PATH, encoding="utf-8") as f:
            catalog_list = json.load(f)
        catalog = {e.get("slug", ""): e for e in catalog_list}

    processed = 0
    skipped = 0
    failed = 0

    for file_path in doc_files:
        slug = file_path.stem
        text_path = TEXT_DIR / f"{slug}.txt"

        # Skip if already extracted (unless --force)
        if not force and text_path.exists() and text_path.stat().st_size > 0:
            skipped += 1
            continue

        log.info(f"Extracting: {file_path.name}")
        result = extract_text(file_path)

        if result["text"]:
            # Clean text for embeddings
            with open(text_path, "w", encoding="utf-8") as f:
                f.write(result["text"])

            # Page-level data for citations
            if result.get("pages"):
                pages_path = TEXT_DIR / f"{slug}.pages.json"
                with open(pages_path, "w", encoding="utf-8") as f:
                    json.dump(result["pages"], f, ensure_ascii=False)

            # Update catalog entry
            if slug in catalog:
                catalog[slug]["page_count"] = result["page_count"]
                catalog[slug]["word_count"] = result["word_count"]
                catalog[slug]["text_extracted"] = True
                if result["pdf_metadata"]:
                    catalog[slug]["pdf_metadata"] = result["pdf_metadata"]

            processed += 1
        else:
            log.warning(f"No text extracted from {file_path.name}")
            failed += 1

    # Save updated catalog
    if CATALOG_PATH.exists():
        with open(CATALOG_PATH, "w", encoding="utf-8") as f:
            json.dump(list(catalog.values()), f, indent=2, ensure_ascii=False)

    log.info(f"Done: {processed} extracted, {skipped} already done, {failed} failed.")


if __name__ == "__main__":
    import sys
    force = "--force" in sys.argv
    extract_all_texts(force=force)
